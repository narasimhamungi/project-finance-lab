"""Structure of the committed workbook. Its values are tested in LibreOffice by excel/regression (not in CI)."""
import re

import pytest
from openpyxl import load_workbook

from conftest import ROOT
from projectfinancelab import regression as reg

WB = ROOT / "excel" / "project_finance_model.xlsx"


@pytest.fixture(scope="module")
def wb():
    return load_workbook(WB)


def test_sheets(wb):
    assert wb.sheetnames == ["README", "Inputs", "Construction", "Operations", "Analysis", "Checks"]


def test_input_cells_are_labelled_values(wb):
    for field, cell in reg.INPUT_CELLS.items():
        c = wb["Inputs"][cell]
        assert c.value is not None and not str(c.value).startswith("="), field
        assert wb["Inputs"][f"B{cell[1:]}"].value, field


def test_output_cells_are_formulas(wb):
    for field, ref, _ in reg.OUTPUTS:
        sheet, cell = ref.split("!")
        assert str(wb[sheet.strip("'")][cell].value).startswith("="), field


def test_sizing_never_reads_the_stress_cells(wb):
    """The frozen-debt design, structurally: nothing on Construction or Operations may reference Inputs rows 26-29."""
    stress = re.compile(r"Inputs!\$?C\$?(26|27|28|29)\b")
    for name in ("Construction", "Operations"):
        for row in wb[name].iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and stress.search(c.value)), f"{name}!{c.coordinate}"


def test_no_hard_coded_numbers_in_time_rows(wb):
    """Operations and Analysis time rows (columns D onward, below the year header) hold formulas only."""
    for name in ("Operations", "Analysis"):
        for row in wb[name].iter_rows(min_row=5, max_row=24, min_col=4):
            for c in row:
                assert c.value is None or str(c.value).startswith("="), f"{name}!{c.coordinate}"
