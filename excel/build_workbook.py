"""Build excel/project_finance_model.xlsx: the project finance model in formulas.

Run from the repo root:  python excel/build_workbook.py   (needs openpyxl; Excel recalculates on open)
Inputs are the shipped hypothetical config with its "downside" case active. Cell map: src/projectfinancelab/regression.py.
Debt and total uses are circular (IDC and the fee depend on the debt). The workbook takes them as pasted values on
Inputs and Checks confirms its own formulas return the same numbers; Excel's iterative calculation is not used.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from projectfinancelab.model import ProjectInputs, Stress, analyse, size_debt  # noqa: E402
from projectfinancelab.regression import INPUT_CELLS, PASS_TEXT  # noqa: E402

CFG = json.loads((ROOT / "examples" / "project_hypothetical.json").read_text(encoding="utf-8"))
ACTIVE = "downside"
OUT = ROOT / "excel" / "project_finance_model.xlsx"
NOPS = 30
COLS = [get_column_letter(4 + k) for k in range(NOPS)]          # D..AG = operating years 1..30
FIRST, LAST = COLS[0], COLS[-1]
EQ_COLS = [get_column_letter(4 + k) for k in range(3 + NOPS)]   # D..AJ = 3 construction slots + 30 operating years

F = "Arial"
BLUE, BLACK, GREEN = Font(name=F, color="0000FF"), Font(name=F), Font(name=F, color="008000")
BOLD, TITLE = Font(name=F, bold=True), Font(name=F, bold=True, size=14)
YELLOW = PatternFill("solid", fgColor="FFFF00")
NUM, PCT, X, INT, DF = "#,##0.00;(#,##0.00);-", "0.00%", '0.0000"x"', "0", "0.000000"


def put(ws, ref, value, font=None, fmt=None, fill=None):
    c = ws[ref]
    c.value = value
    c.font = font or (GREEN if isinstance(value, str) and value.startswith("=") and "!" in value else BLACK)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill


def widths(ws, first=40, rest=12, label_col="B"):
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions[label_col].width = first
    for col in range(3, 40):
        ws.column_dimensions[get_column_letter(col)].width = rest


p = ProjectInputs.from_dict(CFG["project"])
stress = Stress(**CFG["cases"][ACTIVE])
fin = size_debt(p)
sponsor, active = analyse(p, fin), analyse(p, fin, stress)

wb = Workbook()
# ------------------------------------------------------------------ README
rd = wb.active
rd.title = "README"
for i, (text, font) in enumerate([
    ("Project finance model (formula version)", TITLE),
    ("Hypothetical inputs. Not a real project, not market terms, not advice, and not any rating agency's model or criteria.", BOLD),
    ("", None),
    ("What it does", BOLD),
    ("Sizes senior debt on the sponsor case as min(PV of CFADS over the tenor / target DSCR, gearing cap x total uses),", None),
    ("sculpts repayments to the resulting DSCR, then runs the active stress on Inputs against that frozen schedule:", None),
    ("DSCR, LLCR, PLCR, lock-up and default-level tests, payment shortfalls, distributions and equity IRR.", None),
    ("It reproduces src/projectfinancelab/model.py; Checks compares the two.", None),
    ("", None),
    ("How to use", BOLD),
    ("Edit only blue cells on Inputs. Stress cells (rows 26-29) change Analysis only; sizing never sees the stress.", None),
    ("Debt and total uses (rows 32-33) are pasted values: IDC and the fee depend on the debt, so the model is circular.", None),
    ("After changing any sponsor input, paste Checks!C5 into Inputs!C32 and Checks!C6 into Inputs!C33 and repeat", None),
    ("until both 'converged' checks are TRUE (a few rounds), or run project-finance-lab and paste its values.", None),
    ("Excel's iterative calculation is deliberately not used: it hides non-convergence.", None),
    ("", None),
    ("Sheets", BOLD),
    ("Inputs; Construction (capex, fee, IDC, pro-rata funding); Operations (sponsor case, sizing, frozen schedule);", None),
    ("Analysis (stressed case on the frozen schedule); Checks (identities and the Python cross-check).", None),
    ("Method and every formula: docs/model_spec.md. Limitations: docs/limitations.md.", None),
], start=1):
    put(rd, f"A{i}", text, font or BLACK)
rd.column_dimensions["A"].width = 120

# ------------------------------------------------------------------ Inputs
inp = wb.create_sheet("Inputs")
put(inp, "A1", "Inputs (edit blue cells only; all hypothetical)", TITLE)
put(inp, "B3", "Sponsor case and financing terms", BOLD)
prof = list(p.capex_profile) + [0.0] * (3 - p.construction_years)
rows = [("construction_years", "Construction years (1-3)", p.construction_years, INT),
        ("capex", "Capex, before IDC and fees (m)", p.capex, NUM),
        ("capex_y1", "Capex share, construction year 1", prof[0], PCT),
        ("capex_y2", "Capex share, construction year 2", prof[1], PCT),
        ("capex_y3", "Capex share, construction year 3", prof[2], PCT),
        ("operation_years", "Operating years (max 30)", p.operation_years, INT),
        ("revenue_y1", "Revenue, operating year 1 (m)", p.revenue_y1, NUM),
        ("revenue_growth", "Revenue growth a year", p.revenue_growth, PCT),
        ("opex_y1", "Opex, operating year 1 (m)", p.opex_y1, NUM),
        ("opex_growth", "Opex growth a year", p.opex_growth, PCT),
        ("tax_rate", "Tax rate (on EBITDA - depreciation, floor 0)", p.tax_rate, PCT),
        ("depreciation_years", "Depreciation years (straight line on capex)", p.depreciation_years, INT),
        ("debt_rate", "Debt all-in fixed rate", p.debt_rate, PCT),
        ("debt_tenor", "Debt tenor, operating years (< operating years)", p.debt_tenor, INT),
        ("target_dscr", "Target DSCR for sizing", p.target_dscr, X),
        ("max_gearing", "Maximum gearing (debt / total uses)", p.max_gearing, PCT),
        ("upfront_fee", "Upfront fee (share of debt)", p.upfront_fee, PCT),
        ("lockup_dscr", "Lock-up DSCR", p.lockup_dscr, X),
        ("default_dscr", "Default-level DSCR", p.default_dscr, X),
        ("revenue_haircut", "Stress: revenue haircut", stress.revenue_haircut, PCT),
        ("opex_increase", "Stress: opex increase", stress.opex_increase, PCT),
        ("delay_months", "Stress: COD delay, months (revenue lost in year 1)", stress.delay_months, "0.0"),
        ("capex_overrun", "Stress: capex overrun, equity-funded (share of capex)", stress.capex_overrun, PCT),
        ("debt_pasted", "Debt (pasted; see README)", fin.debt, NUM),
        ("uses_pasted", "Total uses (pasted; see README)", fin.total_uses, NUM)]
for field, label, value, fmt in rows:
    r = INPUT_CELLS[field][1:]
    put(inp, f"B{r}", label)
    put(inp, f"C{r}", value, BLUE, fmt, YELLOW if field in ("target_dscr", "max_gearing", "debt_pasted", "uses_pasted") else None)
put(inp, "B25", f"Active stress (shipped: '{ACTIVE}' case; stress never changes the debt)", BOLD)
put(inp, "B31", "Circular values, pasted", BOLD)
put(inp, "D32", "Paste from Checks!C5 once converged")
put(inp, "D33", "Paste from Checks!C6 once converged")
put(inp, "F4", "Legend", BOLD)
put(inp, "F5", "Blue text = input", BLUE)
put(inp, "F6", "Yellow fill = key assumption or pasted value", BLACK, None, YELLOW)
put(inp, "F7", "Source: examples/project_hypothetical.json (hypothetical)")
put(inp, "F9", "Shipped cases (illustrative layering, not any agency's rating case)", BOLD)
for i, (name, s) in enumerate(CFG["cases"].items()):
    put(inp, f"F{10 + i}", f"{name}: " + (", ".join(f"{k} {v}" for k, v in s.items()) or "no stress"))
widths(inp, 52, 14)
inp.column_dimensions["D"].width = 26
inp.column_dimensions["F"].width = 60

# ------------------------------------------------------------------ Construction
con = wb.create_sheet("Construction")
put(con, "A1", "Construction: capex, fee, IDC and pro-rata funding", TITLE)
labels = {4: "Construction year", 5: "Active", 6: "Capex", 7: "Upfront fee", 8: "Opening debt balance",
          9: "IDC (rate x opening balance, capitalised)", 10: "Uses", 11: "Debt draw (debt share x uses)",
          12: "Closing debt balance", 13: "Equity"}
for r, t in labels.items():
    put(con, f"B{r}", t, BOLD if r == 4 else BLACK)
for i, c in enumerate("CDE"):
    prev = "CDE"[i - 1] if i else None
    put(con, f"{c}4", i + 1, BOLD, INT)
    put(con, f"{c}5", f"=IF({c}$4<=Inputs!$C$5,1,0)", None, INT)
    put(con, f"{c}6", f"=Inputs!$C$6*INDEX(Inputs!$C$7:$C$9,{c}$4)*{c}5", None, NUM)
    put(con, f"{c}7", f"=IF({c}$4=1,Inputs!$C$21*Inputs!$C$32,0)", None, NUM)
    put(con, f"{c}8", f"={prev}12" if prev else 0, None, NUM)
    put(con, f"{c}9", f"=Inputs!$C$17*{c}8*{c}5", None, NUM)
    put(con, f"{c}10", f"={c}6+{c}7+{c}9", None, NUM)
    put(con, f"{c}11", f"=$C$16*{c}10", None, NUM)
    put(con, f"{c}12", f"={c}8+{c}11", None, NUM)
    put(con, f"{c}13", f"={c}10-{c}11", None, NUM)
for r, t, f, fmt in [(15, "Total uses (calculated)", "=SUM(C10:E10)", NUM),
                     (16, "Debt share (pasted debt / pasted uses)", "=Inputs!C32/Inputs!C33", PCT),
                     (17, "Debt balance at COD", "=E12", NUM),
                     (18, "Total equity", "=SUM(C13:E13)", NUM)]:
    put(con, f"B{r}", t, BOLD)
    put(con, f"C{r}", f, None, fmt)
widths(con, 44, 14)

# ------------------------------------------------------------------ Operations (sponsor case, sizing, frozen schedule)
ops = wb.create_sheet("Operations")
put(ops, "A1", "Operations: sponsor case, debt sizing and the frozen repayment schedule", TITLE)
olabels = {4: "Operating year", 5: "Operating", 6: "In debt tenor", 7: "Revenue", 8: "Opex", 9: "EBITDA",
           10: "Depreciation", 11: "Tax", 12: "CFADS (sponsor case)", 13: "Discount factor, end of year",
           14: "Discount factor, start of year", 16: "Debt service (sculpted; frozen)", 17: "Opening balance",
           18: "Interest", 19: "Principal", 20: "Closing balance", 22: "LLCR, sponsor case (must equal sculpting DSCR)"}
for r, t in olabels.items():
    put(ops, f"B{r}", t, BOLD if r in (4, 12, 16) else BLACK)
for k, c in enumerate(COLS):
    prev = COLS[k - 1] if k else None
    put(ops, f"{c}4", k + 1, BOLD, INT)
    put(ops, f"{c}5", f"=IF({c}$4<=Inputs!$C$10,1,0)", None, INT)
    put(ops, f"{c}6", f"=IF({c}$4<=Inputs!$C$18,1,0)*{c}5", None, INT)
    put(ops, f"{c}7", f"=Inputs!$C$11*(1+Inputs!$C$12)^({c}$4-1)*{c}5", None, NUM)
    put(ops, f"{c}8", f"=Inputs!$C$13*(1+Inputs!$C$14)^({c}$4-1)*{c}5", None, NUM)
    put(ops, f"{c}9", f"={c}7-{c}8", None, NUM)
    put(ops, f"{c}10", f"=IF({c}$4<=Inputs!$C$16,Inputs!$C$6/Inputs!$C$16,0)*{c}5", None, NUM)
    put(ops, f"{c}11", f"=Inputs!$C$15*MAX(0,{c}9-{c}10)", None, NUM)
    put(ops, f"{c}12", f"={c}9-{c}11", None, NUM)
    put(ops, f"{c}13", f"=1/(1+Inputs!$C$17)^{c}$4", None, DF)
    put(ops, f"{c}14", f"=1/(1+Inputs!$C$17)^({c}$4-1)", None, DF)
    put(ops, f"{c}16", f"={c}6*{c}12/$C$34", None, NUM)
    put(ops, f"{c}17", f"={prev}20" if prev else "=Construction!$C$17", None, NUM)
    put(ops, f"{c}18", f"=Inputs!$C$17*{c}17*{c}6", None, NUM)
    put(ops, f"{c}19", f"={c}16-{c}18", None, NUM)
    put(ops, f"{c}20", f"={c}17-{c}19", None, NUM)
    put(ops, f"{c}22", f"=IF({c}6=1,SUMPRODUCT({c}12:${LAST}$12,{c}13:${LAST}$13,{c}6:${LAST}$6)/{c}14/{c}17,0)", None, X)
put(ops, "B29", "Sizing", BOLD)
for r, t, f, fmt in [
        (30, "PV at COD of sponsor CFADS over the tenor", f"=SUMPRODUCT({FIRST}12:{LAST}12,{FIRST}13:{LAST}13,{FIRST}6:{LAST}6)", NUM),
        (31, "Debt capacity at target DSCR", "=C30/Inputs!C19", NUM),
        (32, "Gearing cap x total uses", "=Inputs!C20*Construction!C15", NUM),
        (33, "Debt (calculated) = min of the two", "=MIN(C31,C32)", NUM),
        (34, "Sculpting DSCR = PV / pasted debt", "=C30/Inputs!C32", X),
        (35, "Binding constraint", '=IF(C31<=C32,"DSCR","gearing")', None),
        (36, "Balance at maturity", f"=INDEX({FIRST}20:{LAST}20,1,Inputs!C18)", NUM),
        (37, "Sum of principal over the tenor", f"=SUMPRODUCT({FIRST}19:{LAST}19,{FIRST}6:{LAST}6)", NUM),
        (38, "Smallest principal payment in the tenor", f"=_xlfn.MINIFS({FIRST}19:{LAST}19,{FIRST}6:{LAST}6,1)", NUM)]:
    put(ops, f"B{r}", t)
    put(ops, f"C{r}", f, None, fmt)
widths(ops, 46, 11)

# ------------------------------------------------------------------ Analysis (stress on the frozen schedule)
an = wb.create_sheet("Analysis")
put(an, "A1", "Analysis: active stress against the frozen debt schedule", TITLE)
alabels = {4: "Operating year", 5: "Operating", 6: "In debt tenor", 7: "Revenue (stressed)", 8: "Opex (stressed)",
           9: "EBITDA", 10: "Tax", 11: "CFADS (stressed)", 12: "Debt service (frozen, from Operations)", 13: "DSCR",
           14: "Opening balance", 15: "LLCR", 16: "Lock-up (DSCR < lock-up level)",
           17: "Default level (DSCR < default level)", 18: "Cash available (b/f + CFADS - debt service)",
           19: "Payment shortfall", 20: "Cash carried forward", 21: "Distributions",
           24: "Equity cash flows: last 3 construction slots, then operating years"}
for r, t in alabels.items():
    put(an, f"B{r}", t, BOLD if r in (4, 11, 13) else BLACK)
put(an, "C20", 0, BLACK, NUM)
for k, c in enumerate(COLS):
    prev = COLS[k - 1] if k else "C"
    put(an, f"{c}4", k + 1, BOLD, INT)
    put(an, f"{c}5", f"=Operations!{c}5", None, INT)
    put(an, f"{c}6", f"=Operations!{c}6", None, INT)
    put(an, f"{c}7", f"=Operations!{c}7*(1-Inputs!$C$26)*IF({c}$4=1,1-Inputs!$C$28/12,1)", None, NUM)
    put(an, f"{c}8", f"=Operations!{c}8*(1+Inputs!$C$27)", None, NUM)
    put(an, f"{c}9", f"={c}7-{c}8", None, NUM)
    put(an, f"{c}10", f"=Inputs!$C$15*MAX(0,{c}9-Operations!{c}10)", None, NUM)
    put(an, f"{c}11", f"={c}9-{c}10", None, NUM)
    put(an, f"{c}12", f"=Operations!{c}16", None, NUM)
    put(an, f"{c}13", f'=IF({c}6=1,{c}11/{c}12,"")', None, X)
    put(an, f"{c}14", f"=Operations!{c}17", None, NUM)
    put(an, f"{c}15", f'=IF({c}6=1,SUMPRODUCT({c}11:${LAST}$11,Operations!{c}13:${LAST}$13,{c}6:${LAST}$6)/Operations!{c}14/{c}14,"")', None, X)
    put(an, f"{c}16", f"=IF({c}6=1,IF({c}13<Inputs!$C$22,1,0),0)", None, INT)
    put(an, f"{c}17", f"=IF({c}6=1,IF({c}13<Inputs!$C$23,1,0),0)", None, INT)
    put(an, f"{c}18", f"={prev}20+{c}11-{c}12", None, NUM)
    put(an, f"{c}19", f"=MAX(0,-{c}18)", None, NUM)
    put(an, f"{c}20", f"=MAX(0,{c}18)-{c}21", None, NUM)
    put(an, f"{c}21", f"=IF({c}16=1,0,MAX(0,{c}18))", None, NUM)
for o, c in enumerate(EQ_COLS[:3]):          # construction years Nc-2, Nc-1, Nc (leading zeros if Nc < 3)
    y = f"(Inputs!$C$5-2+{o})"
    put(an, f"{c}24", f"=IF({y}>=1,-INDEX(Construction!$C$13:$E$13,1,{y})-IF({y}=Inputs!$C$5,Inputs!$C$29*Inputs!$C$6,0),0)",
        None, NUM)
for k, c in enumerate(EQ_COLS[3:]):
    put(an, f"{c}24", f"={COLS[k]}21", None, NUM)
put(an, "B27", "Results", BOLD)
for r, t, f, fmt in [
        (28, "Minimum DSCR", f"=MIN({FIRST}13:{LAST}13)", X),
        (29, "Average DSCR", f"=AVERAGE({FIRST}13:{LAST}13)", X),
        (30, "LLCR at COD", f"={FIRST}15", X),
        (31, "Minimum LLCR", f"=MIN({FIRST}15:{LAST}15)", X),
        (32, "PLCR at COD", f"=SUMPRODUCT({FIRST}11:{LAST}11,Operations!{FIRST}13:{LAST}13,{FIRST}5:{LAST}5)/{FIRST}14", X),
        (33, "Lock-up years", f"=SUM({FIRST}16:{LAST}16)", INT),
        (34, "Default-level years", f"=SUM({FIRST}17:{LAST}17)", INT),
        (35, "Total payment shortfall", f"=SUM({FIRST}19:{LAST}19)", NUM),
        (36, "Total distributions", f"=SUM({FIRST}21:{LAST}21)", NUM),
        (37, "Equity IRR (second guess -10% if the default fails)",
         f'=IFERROR(IRR({EQ_COLS[0]}24:{EQ_COLS[-1]}24),IFERROR(IRR({EQ_COLS[0]}24:{EQ_COLS[-1]}24,-0.1),"n/a"))', PCT)]:
    put(an, f"B{r}", t, BOLD if r in (28, 37) else BLACK)
    put(an, f"C{r}", f, None, fmt)
widths(an, 46, 11)

# ------------------------------------------------------------------ Checks
ck = wb.create_sheet("Checks")
put(ck, "A1", "Checks", TITLE)
put(ck, "A3", "Model identities (must all be TRUE)", BOLD)
put(ck, "C4", "Calculated", BOLD)
checks = [
    ("Debt converged: calculated debt = pasted debt", "=ABS(C5-Inputs!C32)<0.000001", "=Operations!C33"),
    ("Uses converged: calculated uses = pasted uses", "=ABS(C6-Inputs!C33)<0.000001", "=Construction!C15"),
    ("Debt balance at COD = debt", "=ABS(Construction!C17-Inputs!C32)<0.000001", None),
    ("Principal over the tenor sums to the debt", "=ABS(Operations!C37-Inputs!C32)<0.000001", None),
    ("Balance is zero at maturity", "=ABS(Operations!C36)<0.000001", None),
    ("Sources = uses (debt + equity = total uses)", "=ABS(Construction!C17+Construction!C18-Construction!C15)<0.000001", None),
    ("Gearing within the cap", "=Inputs!C32<=Inputs!C20*Construction!C15+0.000001", None),
    ("Sponsor-case LLCR = sculpting DSCR in every tenor year",
     f"=SUMPRODUCT(Operations!{FIRST}6:{LAST}6,ABS(Operations!{FIRST}22:{LAST}22-Operations!C34))<0.000000001", None),
    ("Frozen debt: debt service on Analysis = sizing schedule",
     f"=SUMPRODUCT(ABS(Analysis!{FIRST}12:{LAST}12-Operations!{FIRST}16:{LAST}16))<0.000000001", None),
    ("Capex profile sums to 1 and is fully placed", "=AND(ABS(SUM(Inputs!C7:C9)-1)<0.000000001,ABS(SUM(Construction!C6:E6)-Inputs!C6)<0.000001)", None),
    ("Layout limits: 1-3 construction years, tenor < operating years <= 30",
     "=AND(Inputs!C5>=1,Inputs!C5<=3,Inputs!C18>=1,Inputs!C18<Inputs!C10,Inputs!C10<=30)", None),
]
for i, (label, formula, calc) in enumerate(checks):
    put(ck, f"A{5 + i}", label)
    put(ck, f"B{5 + i}", formula)
    if calc:
        put(ck, f"C{5 + i}", calc, None, NUM)
put(ck, "A17", "Overall", BOLD)
put(ck, "B17", f'=IF(AND(B5:B15),"{PASS_TEXT}","CHECK FAILED")', BOLD)
put(ck, "A19", "Information (not a pass/fail test): smallest principal payment is not negative")
put(ck, "B19", "=Operations!C38>=0")
put(ck, "A21", f"Cross-check against the Python model: shipped inputs, '{ACTIVE}' stress active", BOLD)
put(ck, "A22", "Reference values from project-finance-lab run --config examples/project_hypothetical.json. "
               "They are expected to differ once inputs change; excel/regression tests other inputs.")
for col, head in zip("ABCD", ["Metric", "Python", "Workbook", "Match"]):
    put(ck, f"{col}24", head, BOLD)
xref = [("Debt", fin.debt, "=Operations!C33", NUM), ("Total uses", fin.total_uses, "=Construction!C15", NUM),
        ("Sculpting DSCR", fin.sculpt_dscr, "=Operations!C34", X),
        (f"Minimum DSCR, {ACTIVE} case", active.min_dscr, "=Analysis!C28", X),
        (f"LLCR at COD, {ACTIVE} case", active.llcr_cod, "=Analysis!C30", X),
        (f"PLCR at COD, {ACTIVE} case", active.plcr_cod, "=Analysis!C32", X),
        (f"Lock-up years, {ACTIVE} case", active.lockup_years, "=Analysis!C33", INT),
        (f"Equity IRR, {ACTIVE} case", active.equity_irr, "=Analysis!C37", PCT)]
for i, (label, py, formula, fmt) in enumerate(xref):
    r = 25 + i
    put(ck, f"A{r}", label)
    put(ck, f"B{r}", py, BLUE, fmt)
    put(ck, f"C{r}", formula, None, fmt)
    put(ck, f"D{r}", f'=IF(ABS(C{r}-B{r})<0.000001,"match","MISMATCH")')
ck.column_dimensions["A"].width = 62
for c in "BCD":
    ck.column_dimensions[c].width = 16

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title != "README"
wb.save(OUT)
print(f"wrote {OUT}")
